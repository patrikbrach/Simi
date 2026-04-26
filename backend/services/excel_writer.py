"""Write the result DataFrame to an Excel file with score-based cell highlighting."""
from __future__ import annotations

import io
from datetime import datetime

import pandas as pd
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

_LOW_SCORE_FILL = PatternFill(start_color="FECACA", end_color="FECACA", fill_type="solid")


def build_result_excel(
    df_a: pd.DataFrame,
    match_values: list[str],
    match_scores: list[int],
    threshold: int,
    label: str = "Match_Value",
    match_cities: list[str] | None = None,
) -> bytes:
    result_df = df_a.copy()
    result_df[label] = match_values
    if match_cities is not None:
        result_df["Match_City"] = match_cities
    result_df["Match_Score"] = match_scores

    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        result_df.to_excel(writer, index=False, sheet_name="SimMatch Result")
        ws = writer.sheets["SimMatch Result"]

        score_col_idx = len(result_df.columns)  # 1-based after to_excel
        score_col_letter = get_column_letter(score_col_idx)

        for r in range(2, len(match_scores) + 2):
            ws[f"{score_col_letter}{r}"].number_format = "0"

        for row_num, score in enumerate(match_scores, start=2):  # row 1 = header
            if score < threshold:
                cell = ws[f"{score_col_letter}{row_num}"]
                cell.fill = _LOW_SCORE_FILL

    return buffer.getvalue()


def result_filename() -> str:
    ts = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
    return f"simmatch_result_{ts}.xlsx"
